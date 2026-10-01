"""Render a tailored resume + cover letter to PDF (designed, text-based, ATS-safe) and DOCX.

Layout checks built in:
  - widows: no paragraph, bullet or list may end with a single word / single list item on its own line
  - resume must fit on one page
"""
import asyncio, html, re
from pathlib import Path

e = html.escape

CSS = """
@import url('https://fonts.googleapis.com/css2?family=DM+Sans:wght@300;400;500&family=Plus+Jakarta+Sans:wght@600;700;800&display=swap');
@page { size: Letter; margin: 0.5in 0.6in; }
*{box-sizing:border-box;margin:0;padding:0}
p,li,.section,.row div{text-wrap:pretty}
body{font-family:'DM Sans',sans-serif;font-weight:300;color:#2D3E5A;font-size:10pt;line-height:1.5;-webkit-print-color-adjust:exact}
.header{display:flex;justify-content:space-between;align-items:flex-end;gap:20px;padding-bottom:10px;margin-bottom:8px;border-bottom:2px solid rgba(27,92,232,.22)}
h1{font-family:'Plus Jakarta Sans',sans-serif;font-weight:800;color:#1A2B4A;font-size:28pt;line-height:.95;letter-spacing:-.03em}
.tagline{margin-top:7px;font-size:9.6pt;color:#7A8BA0}
.contact{text-align:right;font-size:9pt;line-height:1.75;color:#7A8BA0}
.contact a{color:#2D3E5A;text-decoration:none}
.label{break-after:avoid;font-family:'Plus Jakarta Sans',sans-serif;font-weight:700;font-size:8pt;letter-spacing:.08em;text-transform:uppercase;color:#1B5CE8;margin:11px 0 5px}
.section{border:1px solid rgba(27,92,232,.13);border-radius:10px;padding:9px 14px 10px;margin-bottom:6px;break-inside:avoid}
.section.flow{break-inside:auto}
.job{margin-bottom:7px;break-inside:avoid}.job:last-child{margin-bottom:0}
.jh{display:flex;justify-content:space-between;align-items:baseline;gap:10px}
.jt{font-family:'Plus Jakarta Sans',sans-serif;font-weight:700;color:#1A2B4A;font-size:10.6pt}
.jc{color:#7A8BA0;font-weight:400;font-size:9.4pt;margin-left:2px}
.jd{color:#7A8BA0;font-style:italic;font-size:9pt;white-space:nowrap}
ul{list-style:none;margin-top:4px}
li{padding-left:12px;text-indent:-12px;margin-bottom:1px}
li:before{content:'';display:inline-block;vertical-align:middle;margin:-2px 8px 0 0;width:4px;height:4px;border-radius:50%;background:rgba(27,92,232,.5)}
.row{display:flex;gap:12px;margin-bottom:3px}.row:last-child{margin-bottom:0}
.list{text-wrap:balance}.it{white-space:nowrap}.sep{color:#9AA8BC}
.k{font-family:'Plus Jakarta Sans',sans-serif;font-weight:700;color:#1A2B4A;font-size:8.8pt;min-width:118px;padding-top:1px}
.meta{display:flex;justify-content:space-between;font-size:9.6pt;color:#7A8BA0;margin:4px 0 16px}
.re{font-family:'Plus Jakarta Sans',sans-serif;font-weight:600;color:#1B5CE8}
.letter p{font-size:11pt;line-height:1.72;margin-bottom:12px}
.sig{font-family:'Plus Jakarta Sans',sans-serif;font-weight:700;color:#1A2B4A;font-size:13pt;margin-top:18px}
"""

WIDOW_JS = r"""() => {
  const out = [];
  document.querySelectorAll('p, li').forEach(el => {
    const words = el.textContent.trim().split(/\s+/); if (words.length < 3) return;
    const orig = el.innerHTML;
    el.innerHTML = words.map(w => `<span>${w}</span>`).join(' ');
    const tops = [...el.querySelectorAll('span')].map(s => Math.round(s.getBoundingClientRect().top));
    const last = tops[tops.length - 1];
    const lines = new Set(tops.map(t => Math.round(t / 4))).size;
    if (lines > 1 && tops.filter(t => Math.abs(t - last) < 3).length < 2) out.push(el.textContent.trim());
    el.innerHTML = orig;
  });
  document.querySelectorAll('.list').forEach(el => {
    const tops = [...el.querySelectorAll('.it')].map(s => Math.round(s.getBoundingClientRect().top));
    const last = tops[tops.length - 1];
    if (new Set(tops).size > 1 && tops.filter(t => t === last).length < 2) out.push(el.textContent.trim());
  });
  return out;
}"""


def items(parts):
    return '<div class="list">' + ''.join(
        f'<span class="it">{e(p)}{"<span class=sep> · </span>" if i < len(parts) - 1 else ""}</span> '
        for i, p in enumerate(parts)) + '</div>'


def header(d):
    p = d["person"]
    return f"""<div class="header"><div><h1>{e(p['name'])}</h1><div class="tagline">{e(d['tagline'])}</div></div>
<div class="contact"><a href="https://{e(p['site'])}" target="_blank">{e(p['site'])}</a><br><a href="mailto:{e(p['email'])}">{e(p['email'])}</a><br>
<a href="{e(p['linkedin_url'])}">{e(p['linkedin'])}</a></div></div>"""


def resume_html(d):
    jobs = "".join(f"""<div class="job"><div class="jh"><div><span class="jt">{e(j['title'])}</span> <span class="jc">· {e(j['company'])}</span></div><span class="jd">{e(j['dates'])}</span></div>
<ul>{''.join(f'<li>{e(b)}</li>' for b in j['bullets'])}</ul></div>""" for j in d["jobs"] if j["bullets"])
    skills = "".join(f'<div class="row"><div class="k">{e(k)}</div>{items(v)}</div>' for k, v in d["skills"] if v)
    return f"""<html><head><meta charset="utf-8"><title>{e(d['person']['name'])} Resume</title><style>{CSS}</style></head><body>{header(d)}
<div class="label">Profile</div><div class="section"><p>{e(d['profile'])}</p></div>
<div class="label">Experience</div><div class="section flow">{jobs}</div>
<div class="label">Skills</div><div class="section">{skills}</div>
<div class="label">Brands, Recognition &amp; Education</div><div class="section">
<div class="row"><div class="k">Selected brands</div>{items(d['brands'])}</div>
<div class="row"><div class="k">Awards</div>{items(d['awards'])}</div>
<div class="row"><div class="k">Education</div><div>{e(d['education'])}</div></div></div></body></html>"""


def letter_html(d):
    ps = "".join(f"<p>{e(p)}</p>" for p in d["letter"])
    return f"""<html><head><meta charset="utf-8"><title>{e(d['person']['name'])} Cover Letter</title><style>{CSS}</style></head><body>{header(d)}
<div class="meta"><span>{e(d['date'])}</span><span class="re">{e(d['re'])}</span></div>
<div class="letter">{ps}<div class="sig">{e(d['person']['name'])}</div><div style="font-size:9.6pt;color:#1B5CE8">{e(d['person']['email'])}</div></div></body></html>"""


async def _render(pairs):
    from playwright.async_api import async_playwright
    report = {}
    async with async_playwright() as p:
        b = await p.chromium.launch()
        for name, doc, path in pairs:
            # measure at the printed width (Letter 8.5in minus 0.6in margins = 7.3in = 701px)
            pg = await b.new_page(viewport={"width": 701, "height": 1000})
            await pg.emulate_media(media="print")
            await pg.set_content(doc, wait_until="networkidle")
            widows = await pg.evaluate(WIDOW_JS)
            await pg.pdf(path=str(path), format="Letter", print_background=True, prefer_css_page_size=True)
            report[name] = {"widows": widows}
            await pg.close()
        await b.close()
    for name, _, path in pairs:
        try:
            import pypdfium2
            report[name]["pages"] = len(pypdfium2.PdfDocument(str(path)))
        except Exception:
            report[name]["pages"] = None
    return report


def docx_resume(d, path):
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
    p = d["person"]
    doc.add_heading(p["name"], 0)
    doc.add_paragraph(d["tagline"].replace("  ·  ", " | "))
    doc.add_paragraph(f"{p['email']} | {p['site']} | {p['linkedin']} | {p['location_line'].replace('  ·  ', ', ')}")
    doc.add_heading("Profile", 1); doc.add_paragraph(d["profile"])
    doc.add_heading("Experience", 1)
    for j in d["jobs"]:
        if not j["bullets"]: continue
        doc.add_heading(f"{j['title']}, {j['company'].replace('  ·  ', ', ')}  ({j['dates']})", 2)
        for b in j["bullets"]: doc.add_paragraph(b, style="List Bullet")
    doc.add_heading("Skills", 1)
    for k, v in d["skills"]:
        if v: doc.add_paragraph(f"{k}: {', '.join(v)}")
    doc.add_heading("Brands, Recognition & Education", 1)
    doc.add_paragraph("Selected brands: " + ", ".join(d["brands"]))
    doc.add_paragraph("Awards: " + ", ".join(d["awards"]))
    doc.add_paragraph("Education: " + d["education"])
    doc.save(str(path))


def docx_letter(d, path):
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    doc.add_paragraph(d["person"]["name"]); doc.add_paragraph(d["person"]["email"] + " | " + d["person"]["site"])
    doc.add_paragraph(d["date"]); doc.add_paragraph(d["re"])
    for para in d["letter"]: doc.add_paragraph(para)
    doc.add_paragraph(d["person"]["name"])
    doc.save(str(path))


def render_all(d, outdir, stem):
    """Render resume + letter. Returns layout report {resume:{widows,pages}, letter:{...}}."""
    outdir = Path(outdir); outdir.mkdir(parents=True, exist_ok=True)
    rh, lh = resume_html(d), letter_html(d)
    (outdir / f"{stem} - Resume.html").write_text(rh)
    (outdir / f"{stem} - Cover Letter.html").write_text(lh)
    rep = asyncio.run(_render([("resume", rh, outdir / f"{stem} - Resume.pdf"),
                               ("letter", lh, outdir / f"{stem} - Cover Letter.pdf")]))
    return rep
